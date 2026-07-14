"""
Real strategy document parsing - .docx via python-docx, .pdf via pypdf.
No mocking: this actually opens and reads the uploaded file.
"""
import io
from docx import Document
from pypdf import PdfReader


def parse_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    # also pull table content, since strategy docs often have tables (personas, pillars, etc.)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())
    return "\n".join(parts)


def parse_strategy_doc(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".docx"):
        return parse_docx(file_bytes)
    elif lower.endswith(".pdf"):
        return parse_pdf(file_bytes)
    elif lower.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type - please upload .docx, .pdf, or .txt")
